"""
SPED Fiscal - Gerenciador de Historico C100/D100  v2.0
=======================================================
Modos de operacao:
  1. INCREMENTAR  - Adiciona um SPED mensal ao historico existente
  2. CRIAR NOVO   - Cria historico do zero a partir de um unico SPED
  3. CRIAR EM LOTE - Cria historico do zero processando uma pasta inteira de SPEDs

Regras em todos os modos:
  - Apenas C100 e D100 sao extraidos
  - Nenhuma linha existente e removida
  - Sem notacao cientifica (DOCUMENTO, CHAVE, DATA EMISSAO, CNPJ gravados como texto)
  - Prevencao contra reprocessamento do mesmo mes
  - Deduplicacao somente quando TODOS os campos sao identicos
  - CNPJ extraido automaticamente do registro 0000
  - Coluna CNPJ adicionada/mantida automaticamente

Dependencias: openpyxl  (pip install openpyxl)
"""

import os
import datetime
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import openpyxl
except ImportError:
    openpyxl = None


# ---------------------------------------------------------------------------
# Mapeamento de campos SPED
# ---------------------------------------------------------------------------
REG0000_DT_INI = 4
REG0000_CNPJ   = 7
C100_NUM_DOC = 8;  C100_CHV_NFE = 9;  C100_DT_DOC = 10
D100_NUM_DOC = 9;  D100_CHV_CTE = 10; D100_DT_DOC = 11
ENCODING_SPED = "latin-1"

COL_REGISTRO  = "REGISTRO"
COL_DOCUMENTO = "DOCUMENTO"
COL_CHAVE     = "CHAVE"
COL_DATA      = "DATA EMISSAO"    # chave interna
COL_DATA_XL   = "DATA EMISSÃO"   # como fica no Excel
COL_COMP      = "CONSTA SPED FISCAL"
COL_CNPJ      = "CNPJ"

HEADER_PADRAO = [COL_REGISTRO, COL_DOCUMENTO, COL_CHAVE, COL_DATA_XL, COL_COMP, COL_CNPJ]


# ---------------------------------------------------------------------------
# Funcoes auxiliares
# ---------------------------------------------------------------------------

def _get(partes, idx):
    return partes[idx].strip() if len(partes) > idx else ""

def _col_key(col):
    if col in ("DATA EMISSÃO", "DATA EMISSAO"): return COL_DATA
    return col

MESES_PT = ["jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez"]

# Formato de exibição da coluna CONSTA SPED FISCAL no Excel.
# O código [$-416] força o locale pt-BR: exibe "dez/24", "abr/26", etc.
FMT_COMP = "[$-416]mmm/yy"

def _fmt_comp(dt):
    """Formata datetime para exibição no log (ex: dez/24)."""
    if isinstance(dt, datetime.datetime):
        return f"{MESES_PT[dt.month-1]}/{str(dt.year)[2:]}"
    return str(dt) if dt else ""

def _parse_comp(val):
    """
    Lê o valor da coluna CONSTA SPED FISCAL do histórico existente.
    Aceita datetime (qualquer histórico) e retorna datetime normalizado (dia=1, sem hora).
    """
    if isinstance(val, datetime.datetime):
        return datetime.datetime(val.year, val.month, 1)
    return None

def _reg_val(reg, col):
    return reg.get(_col_key(col), "") or ""


# ---------------------------------------------------------------------------
# Logica principal
# ---------------------------------------------------------------------------

def parsear_sped(caminho_sped):
    """Le o SPED TXT e extrai C100 e D100. Retorna (competencia, cnpj, registros)."""
    competencia = None; cnpj = ""; registros = []
    with open(caminho_sped, "r", encoding=ENCODING_SPED, errors="replace") as f:
        for linha in f:
            linha = linha.rstrip("\n\r")
            if not linha: continue
            partes = linha.split("|")
            if len(partes) < 2: continue
            tipo = partes[1].strip().upper()

            if tipo == "0000":
                try:
                    s = partes[REG0000_DT_INI].strip()
                    competencia = datetime.datetime(int(s[4:8]), int(s[2:4]), int(s[0:2]))
                except: pass
                try: cnpj = partes[REG0000_CNPJ].strip()
                except: pass
            elif tipo in ("C100", "D100"):
                nd = C100_NUM_DOC if tipo=="C100" else D100_NUM_DOC
                nc = C100_CHV_NFE if tipo=="C100" else D100_CHV_CTE
                nt = C100_DT_DOC  if tipo=="C100" else D100_DT_DOC
                registros.append({
                    COL_REGISTRO: tipo, COL_DOCUMENTO: _get(partes, nd),
                    COL_CHAVE: _get(partes, nc), COL_DATA: _get(partes, nt),
                    COL_COMP: None, COL_CNPJ: None,
                })
    for r in registros:
        r[COL_COMP] = competencia
        r[COL_CNPJ] = cnpj
    return competencia, cnpj, registros


def carregar_historico(caminho_xlsx):
    """Abre historico existente. Retorna (wb, ws, comps_existentes, linhas_existentes, header)."""
    wb = openpyxl.load_workbook(caminho_xlsx)
    ws = wb.active
    header = []; comps = set(); linhas = set()
    primeira = True
    for row in ws.iter_rows(values_only=True):
        if primeira:
            header = [str(c).strip() if c else "" for c in row]
            primeira = False; continue
        if not any(c is not None for c in row): continue
        comp_parsed = _parse_comp(row[4])
        if comp_parsed: comps.add(comp_parsed)
        linhas.add(tuple(str(c) if c is not None else "" for c in row))
    return wb, ws, comps, linhas, header


def criar_historico_vazio(caminho_saida):
    """Cria um novo workbook com o header padrao. Retorna (wb, ws)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CONSOLIDADO"
    ws.append(HEADER_PADRAO)
    # Formatar colunas: CONSTA SPED FISCAL como data pt-BR, demais como texto
    for col_idx, col in enumerate(HEADER_PADRAO, 1):
        fmt = FMT_COMP if col == COL_COMP else "@"
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(col_idx)
        ].number_format = fmt
    wb.save(caminho_saida)
    return wb, ws


def _escrever_registros(ws, header_escrita, registros_a_adicionar, proxima_linha):
    """Grava lista de registros no ws a partir de proxima_linha."""
    for reg in registros_a_adicionar:
        for ci, col in enumerate(header_escrita, 1):
            cell = ws.cell(row=proxima_linha, column=ci)
            if col == COL_COMP:
                # Gravar como data real (dia=1, sem hora) com formato pt-BR
                comp_dt = reg.get(COL_COMP)
                if isinstance(comp_dt, datetime.datetime):
                    cell.value = datetime.datetime(comp_dt.year, comp_dt.month, 1)
                else:
                    cell.value = comp_dt
                cell.number_format = FMT_COMP
            else:
                cell.number_format = "@"
                cell.value = str(_reg_val(reg, col))
        proxima_linha += 1
    return proxima_linha


def _proxima_linha_livre(ws, ncols):
    """Encontra a proxima linha vazia percorrendo de baixo para cima."""
    for ri in range(ws.max_row, 0, -1):
        if any(ws.cell(row=ri, column=c).value is not None for c in range(1, ncols + 1)):
            return ri + 1
    return 2  # so header


def _normalizar(reg, header):
    return tuple(str(_reg_val(reg, col)) for col in header)


# ---------------------------------------------------------------------------
# Modo 1: Incrementar historico existente (um SPED)
# ---------------------------------------------------------------------------

def incrementar_historico(caminho_historico, caminho_sped, caminho_saida, callback_log=None):
    """
    Adiciona registros de um SPED ao historico existente.
    Retorna (n_adicionados, n_ignorados, aviso_reprocessamento).
    """
    def log(m):
        if callback_log: callback_log(m)

    log("Lendo arquivo SPED...")
    comp, cnpj, regs_novos = parsear_sped(caminho_sped)
    if comp is None:
        raise ValueError("Registro 0000 nao encontrado no SPED.")
    log(f"  Competencia  : {comp.strftime('%m/%Y')}")
    log(f"  CNPJ         : {cnpj}")
    log(f"  C100 + D100  : {len(regs_novos)} registros")

    log("Carregando historico acumulado...")
    wb, ws, comps_exist, linhas_exist, header_orig = carregar_historico(caminho_historico)
    log(f"  Colunas      : {header_orig}")
    log(f"  Linhas       : {len(linhas_exist)}")
    log(f"  Competencias : {sorted(_fmt_comp(c) for c in comps_exist)}")

    aviso = comp in comps_exist
    if aviso:
        log(f"  ATENCAO: {comp.strftime('%m/%Y')} ja existe! Aplicando deduplicacao estrita.")

    # Deduplicar com header ORIGINAL (antes de adicionar CNPJ)
    n_dup = 0; to_add = []
    for reg in regs_novos:
        t = _normalizar(reg, header_orig)
        if t in linhas_exist: n_dup += 1
        else: to_add.append(reg); linhas_exist.add(t)
    log(f"  A adicionar  : {len(to_add)}")
    log(f"  Duplicatas   : {n_dup}")

    # Adicionar coluna CNPJ APOS deduplicacao
    header_escrita = list(header_orig)
    if COL_CNPJ not in header_escrita:
        ws.cell(row=1, column=len(header_escrita) + 1).value = COL_CNPJ
        header_escrita.append(COL_CNPJ)
        log("  Coluna CNPJ adicionada ao historico.")

    # Garantir que a coluna CONSTA SPED FISCAL tenha o formato correto
    # (históricos antigos podem ter datetime com hora ou formato padrão)
    if COL_COMP in header_escrita:
        ci_comp = header_escrita.index(COL_COMP) + 1
        for ri in range(2, ws.max_row + 1):
            cell = ws.cell(row=ri, column=ci_comp)
            if cell.value is not None:
                # Normalizar: remover hora se existir
                if isinstance(cell.value, datetime.datetime) and (
                    cell.value.hour != 0 or cell.value.minute != 0
                ):
                    cell.value = datetime.datetime(cell.value.year, cell.value.month, 1)
                cell.number_format = FMT_COMP

    prox = _proxima_linha_livre(ws, len(header_escrita))
    _escrever_registros(ws, header_escrita, to_add, prox)

    log(f"Salvando em: {caminho_saida}")
    wb.save(caminho_saida)
    log("Concluido!")
    return len(to_add), n_dup, aviso


# ---------------------------------------------------------------------------
# Modo 2: Criar historico a partir de um unico SPED
# ---------------------------------------------------------------------------

def criar_de_sped_unico(caminho_sped, caminho_saida, callback_log=None):
    """
    Cria um historico novo a partir de um unico arquivo SPED.
    Retorna (n_adicionados,).
    """
    def log(m):
        if callback_log: callback_log(m)

    log("Lendo arquivo SPED...")
    comp, cnpj, regs = parsear_sped(caminho_sped)
    if comp is None:
        raise ValueError("Registro 0000 nao encontrado no SPED.")
    log(f"  Competencia  : {comp.strftime('%m/%Y')}")
    log(f"  CNPJ         : {cnpj}")
    log(f"  C100 + D100  : {len(regs)} registros")

    log("Criando novo historico...")
    wb, ws = criar_historico_vazio(caminho_saida)
    _escrever_registros(ws, HEADER_PADRAO, regs, 2)

    log(f"Salvando em: {caminho_saida}")
    wb.save(caminho_saida)
    log(f"Concluido! {len(regs)} registros gravados.")
    return (len(regs),)


# ---------------------------------------------------------------------------
# Modo 3: Criar historico em lote (pasta com varios SPEDs)
# ---------------------------------------------------------------------------

def criar_de_pasta(caminho_pasta, caminho_saida, callback_log=None):
    """
    Cria um historico novo processando todos os TXT de uma pasta,
    em ordem cronologica de competencia.
    Retorna (n_total_adicionados, n_arquivos_processados, n_arquivos_sem_dados).
    """
    def log(m):
        if callback_log: callback_log(m)

    # Listar TXTs
    txts = [f for f in os.listdir(caminho_pasta) if f.lower().endswith(".txt")]
    if not txts:
        raise ValueError(f"Nenhum arquivo .txt encontrado em:\n{caminho_pasta}")
    log(f"Encontrados {len(txts)} arquivos .txt na pasta.")

    # Parsear todos e ordenar por competencia
    log("Lendo e ordenando arquivos por competencia...")
    speds = []
    for fname in txts:
        fpath = os.path.join(caminho_pasta, fname)
        try:
            comp, cnpj, regs = parsear_sped(fpath)
            if comp is None:
                log(f"  IGNORADO (sem registro 0000): {fname}")
                continue
            speds.append((comp, cnpj, regs, fname))
        except Exception as e:
            log(f"  ERRO ao ler {fname}: {e}")

    if not speds:
        raise ValueError("Nenhum arquivo SPED valido encontrado na pasta.")

    speds.sort(key=lambda x: x[0])  # ordem cronologica
    log(f"  {len(speds)} arquivos validos, ordenados por competencia.")

    # Criar historico vazio
    log("Criando novo historico...")
    wb, ws = criar_historico_vazio(caminho_saida)

    # Processar cada SPED acumulando deduplicacao
    linhas_exist = set()
    comps_vistas = set()
    n_total = 0
    n_sem_dados = 0
    prox = 2

    for comp, cnpj, regs, fname in speds:
        comp_str = comp.strftime("%m/%Y")

        if comp in comps_vistas:
            log(f"  ATENCAO: competencia {comp_str} duplicada nos arquivos. "
                f"Aplicando deduplicacao em: {fname}")
        comps_vistas.add(comp)

        to_add = []
        n_dup_local = 0
        for reg in regs:
            t = _normalizar(reg, HEADER_PADRAO)
            if t in linhas_exist:
                n_dup_local += 1
            else:
                to_add.append(reg)
                linhas_exist.add(t)

        if not regs:
            n_sem_dados += 1
            log(f"  [{comp_str}] {fname} — sem C100/D100")
        elif n_dup_local > 0:
            log(f"  [{comp_str}] {fname} — {len(to_add)} adicionados, {n_dup_local} duplicatas ignoradas")
        else:
            log(f"  [{comp_str}] {fname} — {len(to_add)} registros")

        prox = _escrever_registros(ws, HEADER_PADRAO, to_add, prox)
        n_total += len(to_add)

    log(f"Salvando em: {caminho_saida}")
    wb.save(caminho_saida)
    log(f"Concluido! {n_total} registros no historico gerado.")
    return n_total, len(speds), n_sem_dados


# ---------------------------------------------------------------------------
# Interface Grafica
# ---------------------------------------------------------------------------

class App(tk.Tk):

    BG       = "#f0f2f5"
    CARD     = "#ffffff"
    CARD2    = "#f0f2f5"
    ENTRY_BG = "#ffffff"
    ACCENT   = "#0078d4"
    ACCENT2  = "#0078d4"
    FG       = "#1a1a1a"
    FG_DIM   = "#94a3b8"
    BTN_BG   = "#0078d4"
    BTN_FG   = "#ffffff"
    BTN_HOV  = "#005fa3"
    LOG_BG   = "#1e1e1e"
    LOG_FG   = "#d4d4d4"
    SEL_BG   = "#0078d4"

    def __init__(self):
        super().__init__()
        self.title("SPED Fiscal  -  Gerenciador de Historico C100/D100")
        self.resizable(False, False)
        self.configure(bg=self.BG)
        self._modo = tk.IntVar(value=1)
        self._build_ui()
        self._check_openpyxl()
        self._atualizar_modo()

    # ------------------------------------------------------------------
    # Construcao da UI
    # ------------------------------------------------------------------

    def _build_ui(self):
        PAD = 18

        # Cabecalho
        fh = tk.Frame(self, bg=self.BG)
        fh.pack(fill="x", padx=PAD, pady=(PAD, 4))
        tk.Label(fh, text="SPED Fiscal  |  Gerenciador de Historico C100 / D100",
                 font=("Segoe UI", 13, "bold"), fg=self.FG, bg=self.BG).pack(anchor="w")
        tk.Label(fh, text="Cria ou incrementa o historico acumulado de C100 e D100",
                 font=("Segoe UI", 9), fg=self.FG_DIM, bg=self.BG).pack(anchor="w")

        ttk.Separator(self).pack(fill="x", padx=PAD, pady=8)

        # Selecao de modo
        fm = tk.LabelFrame(self, text="  Modo de operacao  ",
                           bg=self.CARD, fg=self.FG_DIM,
                           font=("Segoe UI", 9), bd=1, relief="groove")
        fm.pack(fill="x", padx=PAD, pady=(0, 8))

        modos = [
            (1, "Incrementar historico existente  (um SPED mensal)"),
            (2, "Criar novo historico  (um unico SPED)"),
            (3, "Criar novo historico em lote  (pasta com varios SPEDs)"),
        ]
        for val, txt in modos:
            tk.Radiobutton(
                fm, text=txt, variable=self._modo, value=val,
                command=self._atualizar_modo,
                bg=self.CARD, fg=self.FG, selectcolor=self.CARD2,
                activebackground=self.CARD, activeforeground=self.FG,
                font=("Segoe UI", 9), cursor="hand2",
            ).pack(anchor="w", padx=12, pady=3)

        # Card de campos (varia conforme o modo)
        self._frame_campos = tk.Frame(self, bg=self.CARD, pady=10)
        self._frame_campos.pack(fill="x", padx=PAD, pady=4)

        # --- Campos modo 1: Incrementar ---
        self._f1 = tk.Frame(self._frame_campos, bg=self.CARD)
        self.var_historico = tk.StringVar()
        self.var_sped_inc  = tk.StringVar()
        self.var_saida_inc = tk.StringVar()
        self._file_row(self._f1, "Historico acumulado  (.xlsx):",
                       self.var_historico, self._browse_historico)
        self._file_row(self._f1, "Arquivo SPED do mes  (.txt):",
                       self.var_sped_inc, self._browse_sped_inc)
        self._file_row(self._f1, "Salvar novo historico  (.xlsx):",
                       self.var_saida_inc, self._browse_saida_inc)

        # --- Campos modo 2: Criar de um SPED ---
        self._f2 = tk.Frame(self._frame_campos, bg=self.CARD)
        self.var_sped_novo  = tk.StringVar()
        self.var_saida_novo = tk.StringVar()
        self._file_row(self._f2, "Arquivo SPED  (.txt):",
                       self.var_sped_novo, self._browse_sped_novo)
        self._file_row(self._f2, "Salvar historico  (.xlsx):",
                       self.var_saida_novo, self._browse_saida_novo)

        # --- Campos modo 3: Criar em lote ---
        self._f3 = tk.Frame(self._frame_campos, bg=self.CARD)
        self.var_pasta      = tk.StringVar()
        self.var_saida_lote = tk.StringVar()
        self._folder_row(self._f3, "Pasta com arquivos SPED  (.txt):",
                         self.var_pasta, self._browse_pasta)
        self._file_row(self._f3, "Salvar historico  (.xlsx):",
                       self.var_saida_lote, self._browse_saida_lote)

        # Botao principal
        fb = tk.Frame(self, bg=self.BG)
        fb.pack(fill="x", padx=PAD, pady=10)
        self.btn = tk.Button(
            fb,
            text="  Executar  ",
            font=("Segoe UI", 11, "bold"),
            bg=self.BTN_BG, fg=self.BTN_FG,
            activebackground=self.BTN_HOV, activeforeground="#fff",
            relief="flat", cursor="hand2", pady=10,
            command=self._executar,
        )
        self.btn.pack(fill="x")

        # Progresso
        self.progress = ttk.Progressbar(self, mode="indeterminate")
        self.progress.pack(fill="x", padx=PAD, pady=(0, 6))

        # Log
        fl = tk.Frame(self, bg=self.CARD)
        fl.pack(fill="both", expand=True, padx=PAD, pady=(0, PAD))
        tk.Label(fl, text="Log de execucao",
                 font=("Segoe UI", 9, "bold"),
                 fg=self.FG_DIM, bg=self.CARD).pack(anchor="w", padx=8, pady=(6, 0))

        ft = tk.Frame(fl, bg=self.CARD)
        ft.pack(fill="both", expand=True, padx=8, pady=(4, 8))
        sb = tk.Scrollbar(ft)
        sb.pack(side="right", fill="y")
        self.txt_log = tk.Text(
            ft, height=14, width=84,
            bg=self.LOG_BG, fg=self.LOG_FG,
            font=("Consolas", 9), relief="flat", bd=0,
            state="disabled", wrap="word", yscrollcommand=sb.set,
        )
        self.txt_log.pack(side="left", fill="both", expand=True)
        sb.config(command=self.txt_log.yview)
        self.txt_log.tag_configure("ERR",  foreground="#ef4444")
        self.txt_log.tag_configure("WARN", foreground="#f59e0b")
        self.txt_log.tag_configure("OK",   foreground="#22c55e")
        self.txt_log.tag_configure("INFO", foreground="#a0f0a0")

        self.minsize(720, 640)

    def _file_row(self, parent, label, var, cmd):
        f = tk.Frame(parent, bg=self.CARD)
        f.pack(fill="x", padx=12, pady=4)
        tk.Label(f, text=label, font=("Segoe UI", 9), fg=self.FG_DIM, bg=self.CARD,
                 width=36, anchor="w").pack(side="left")
        tk.Entry(f, textvariable=var, bg=self.ENTRY_BG, fg=self.FG,
                 insertbackground=self.FG, relief="flat",
                 font=("Segoe UI", 9), width=36).pack(side="left", padx=(4, 4))
        tk.Button(f, text="  ...  ", bg=self.ACCENT, fg="#fff",
                  activebackground=self.BTN_HOV, activeforeground="#fff",
                  relief="flat", cursor="hand2", font=("Segoe UI", 9, "bold"),
                  command=cmd).pack(side="left")

    def _folder_row(self, parent, label, var, cmd):
        f = tk.Frame(parent, bg=self.CARD)
        f.pack(fill="x", padx=12, pady=4)
        tk.Label(f, text=label, font=("Segoe UI", 9), fg=self.FG_DIM, bg=self.CARD,
                 width=36, anchor="w").pack(side="left")
        tk.Entry(f, textvariable=var, bg=self.ENTRY_BG, fg=self.FG,
                 insertbackground=self.FG, relief="flat",
                 font=("Segoe UI", 9), width=36).pack(side="left", padx=(4, 4))
        tk.Button(f, text="  ...  ", bg=self.ACCENT2, fg="#fff",
                  activebackground=self.BTN_HOV, activeforeground="#fff",
                  relief="flat", cursor="hand2", font=("Segoe UI", 9, "bold"),
                  command=cmd).pack(side="left")

    # ------------------------------------------------------------------
    # Troca de modo
    # ------------------------------------------------------------------

    def _atualizar_modo(self):
        modo = self._modo.get()
        for f in (self._f1, self._f2, self._f3):
            f.pack_forget()
        if modo == 1:
            self._f1.pack(fill="x")
            self.btn.config(text="  Processar e Incrementar Historico  ")
        elif modo == 2:
            self._f2.pack(fill="x")
            self.btn.config(text="  Criar Historico a partir do SPED  ")
        else:
            self._f3.pack(fill="x")
            self.btn.config(text="  Criar Historico em Lote  ")

    # ------------------------------------------------------------------
    # Browsing
    # ------------------------------------------------------------------

    def _browse_historico(self):
        p = filedialog.askopenfilename(title="Selecionar historico acumulado",
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")])
        if p:
            self.var_historico.set(p)
            if not self.var_saida_inc.get():
                base, ext = os.path.splitext(p)
                self.var_saida_inc.set(base + "_ATUALIZADO" + ext)

    def _browse_sped_inc(self):
        p = filedialog.askopenfilename(title="Selecionar SPED do mes",
            filetypes=[("Texto", "*.txt"), ("Todos", "*.*")])
        if p: self.var_sped_inc.set(p)

    def _browse_saida_inc(self):
        p = filedialog.asksaveasfilename(title="Salvar historico atualizado",
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("Todos", "*.*")])
        if p: self.var_saida_inc.set(p)

    def _browse_sped_novo(self):
        p = filedialog.askopenfilename(title="Selecionar arquivo SPED",
            filetypes=[("Texto", "*.txt"), ("Todos", "*.*")])
        if p:
            self.var_sped_novo.set(p)
            if not self.var_saida_novo.get():
                pasta = os.path.dirname(p)
                self.var_saida_novo.set(os.path.join(pasta, "historico_novo.xlsx"))

    def _browse_saida_novo(self):
        p = filedialog.asksaveasfilename(title="Salvar novo historico",
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("Todos", "*.*")])
        if p: self.var_saida_novo.set(p)

    def _browse_pasta(self):
        p = filedialog.askdirectory(title="Selecionar pasta com arquivos SPED")
        if p:
            self.var_pasta.set(p)
            if not self.var_saida_lote.get():
                self.var_saida_lote.set(os.path.join(p, "historico_lote.xlsx"))

    def _browse_saida_lote(self):
        p = filedialog.asksaveasfilename(title="Salvar historico em lote",
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("Todos", "*.*")])
        if p: self.var_saida_lote.set(p)

    # ------------------------------------------------------------------
    # Log
    # ------------------------------------------------------------------

    def _log(self, msg, tag="INFO"):
        self.txt_log.config(state="normal")
        ts = datetime.datetime.now().strftime("%H:%M:%S")
        self.txt_log.insert("end", f"[{ts}] {msg}\n", tag)
        self.txt_log.see("end")
        self.txt_log.config(state="disabled")
        self.update_idletasks()

    def _check_openpyxl(self):
        if openpyxl is None:
            self._log("ERRO: openpyxl nao instalado. Execute: pip install openpyxl", "ERR")

    # ------------------------------------------------------------------
    # Execucao
    # ------------------------------------------------------------------

    def _executar(self):
        if openpyxl is None:
            messagebox.showerror("Dependencia ausente",
                                 "Instale o openpyxl:\n\npip install openpyxl")
            return

        modo = self._modo.get()

        if modo == 1:
            self._executar_modo1()
        elif modo == 2:
            self._executar_modo2()
        else:
            self._executar_modo3()

    def _validar(self, checks):
        """checks = lista de (condicao_erro, mensagem). Retorna False se houver erro."""
        erros = [msg for cond, msg in checks if cond]
        if erros:
            messagebox.showerror("Erro", "\n\n".join(erros))
            return False
        return True

    def _pre_run(self, label_historico, label_sped_ou_pasta, label_saida):
        self.btn.config(state="disabled")
        self.progress.start(12)
        self._log("-" * 62)
        self._log(f"Modo     : {['', 'Incrementar', 'Criar (unico)', 'Criar (lote)'][self._modo.get()]}")
        if label_historico: self._log(f"Historico: {os.path.basename(label_historico)}")
        self._log(f"Entrada  : {os.path.basename(label_sped_ou_pasta)}")
        self._log(f"Saida    : {os.path.basename(label_saida)}")

    def _pos_run(self, msg_ok, msg_warn=None):
        self.progress.stop()
        self.btn.config(state="normal")
        if msg_warn:
            self._log(msg_warn, "WARN")
        self._log(msg_ok, "OK")
        messagebox.showinfo("Concluido", msg_ok + ("\n\n" + msg_warn if msg_warn else ""))

    def _err_run(self, err):
        self.progress.stop()
        self.btn.config(state="normal")
        self._log(f"ERRO: {err}", "ERR")
        messagebox.showerror("Erro no processamento", err)

    # Modo 1
    def _executar_modo1(self):
        hist = self.var_historico.get().strip()
        sped = self.var_sped_inc.get().strip()
        saida = self.var_saida_inc.get().strip()
        if not self._validar([
            (not hist,  "Selecione o historico acumulado."),
            (not sped,  "Selecione o arquivo SPED do mes."),
            (not saida, "Defina o caminho de saida."),
            (hist and not os.path.isfile(hist),  f"Historico nao encontrado:\n{hist}"),
            (sped and not os.path.isfile(sped),  f"SPED nao encontrado:\n{sped}"),
        ]): return
        self._pre_run(hist, sped, saida)
        def run():
            try:
                n_add, n_dup, aviso = incrementar_historico(
                    hist, sped, saida,
                    callback_log=lambda m: self.after(0, self._log, m))
                warn = (f"ATENCAO: competencia ja existia. Apenas {n_add} registros ineditos foram incluidos."
                        if aviso else None)
                ok = (f"Historico atualizado!\n\n"
                      f"  Registros adicionados :  {n_add}\n"
                      f"  Duplicatas ignoradas  :  {n_dup}\n\n"
                      f"Arquivo salvo em:\n{saida}")
                self.after(0, self._pos_run, ok, warn)
            except Exception as e:
                self.after(0, self._err_run, str(e))
        threading.Thread(target=run, daemon=True).start()

    # Modo 2
    def _executar_modo2(self):
        sped = self.var_sped_novo.get().strip()
        saida = self.var_saida_novo.get().strip()
        if not self._validar([
            (not sped,  "Selecione o arquivo SPED."),
            (not saida, "Defina o caminho de saida."),
            (sped and not os.path.isfile(sped), f"SPED nao encontrado:\n{sped}"),
        ]): return
        self._pre_run(None, sped, saida)
        def run():
            try:
                (n_add,) = criar_de_sped_unico(
                    sped, saida,
                    callback_log=lambda m: self.after(0, self._log, m))
                ok = (f"Historico criado com sucesso!\n\n"
                      f"  Registros gravados :  {n_add}\n\n"
                      f"Arquivo salvo em:\n{saida}")
                self.after(0, self._pos_run, ok)
            except Exception as e:
                self.after(0, self._err_run, str(e))
        threading.Thread(target=run, daemon=True).start()

    # Modo 3
    def _executar_modo3(self):
        pasta = self.var_pasta.get().strip()
        saida = self.var_saida_lote.get().strip()
        if not self._validar([
            (not pasta,  "Selecione a pasta com os arquivos SPED."),
            (not saida,  "Defina o caminho de saida."),
            (pasta and not os.path.isdir(pasta), f"Pasta nao encontrada:\n{pasta}"),
        ]): return
        self._pre_run(None, pasta, saida)
        def run():
            try:
                n_total, n_arq, n_sem = criar_de_pasta(
                    pasta, saida,
                    callback_log=lambda m: self.after(0, self._log, m))
                ok = (f"Historico em lote criado com sucesso!\n\n"
                      f"  Arquivos processados  :  {n_arq}\n"
                      f"  Sem C100/D100         :  {n_sem}\n"
                      f"  Registros no historico:  {n_total}\n\n"
                      f"Arquivo salvo em:\n{saida}")
                self.after(0, self._pos_run, ok)
            except Exception as e:
                self.after(0, self._err_run, str(e))
        threading.Thread(target=run, daemon=True).start()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    app = App()
    app.mainloop()
